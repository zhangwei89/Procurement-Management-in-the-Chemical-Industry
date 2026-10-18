import os, sys, re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

base_dir = r"C:\Users\zdj\Documents\ChatGPT\New project\化工行业采购管理"
output_docx = os.path.join(base_dir, "《化工行业采购管理》_出版送审稿.docx")

chapters_order = [
    "00_选题策划与全书大纲/00_前言与全书导读.md",
    "00_选题策划与全书大纲/01_选题策划案_Book_Proposal.md",
    "00_选题策划与全书大纲/02_全书总大纲与目录结构.md",
    "01_第一篇_采购管理理论与决策基础/第1章_化工企业采购管理概论.md",
    "01_第一篇_采购管理理论与决策基础/第2章_采购管理原理与供应资源配置.md",
    "02_第二篇_采购业务流程与全过程管理/第3章_采购需求与计划管理.md",
    "02_第二篇_采购业务流程与全过程管理/第4章_采购技术管理与技术需求审查.md",
    "02_第二篇_采购业务流程与全过程管理/第5章_供应市场分析与采购寻源.md",
    "02_第二篇_采购业务流程与全过程管理/第6章_采购方式选择与采购实施.md",
    "02_第二篇_采购业务流程与全过程管理/第7章_采购价格与成本管理.md",
    "02_第二篇_采购业务流程与全过程管理/第8章_合同管理与采购履约.md",
    "03_第三篇_化工行业专业采购管理/第9章_化工原辅材料采购管理.md",
    "03_第三篇_化工行业专业采购管理/第10章_设备备件与仪器仪表采购管理.md",
    "03_第三篇_化工行业专业采购管理/第11章_工程项目检维修与技术服务采购.md",
    "04_第四篇_供应链管理与供应保障/第12章_供应商全生命周期管理.md",
    "04_第四篇_供应链管理与供应保障/第13章_库存物流与供应风险管理.md",
    "05_第五篇_信息化与人工智能采购管理/第14章_信息化采购与企业数字化系统管理.md",
    "05_第五篇_信息化与人工智能采购管理/第15章_人工智能产品与服务采购管理.md",
    "05_第五篇_信息化与人工智能采购管理/第16章_化工企业采购数字化与AI应用实践.md",
    "06_第六篇_采购治理与管理提升/第17章_采购监督审计与风险控制.md",
    "06_第六篇_采购治理与管理提升/第18章_采购组织绩效评价与持续改进.md",
    "00_选题策划与全书大纲/99_后记与致谢.md",
    "07_附录与实务工具包/附录01_采购管理制度类_管理办法与授权矩阵.md",
    "07_附录与实务工具包/附录02_业务操作类_规格书与计划表单.md",
    "07_附录与实务工具包/附录03_供应商管理类_准入审核与绩效评价表.md",
    "07_附录与实务工具包/附录04_成本分析类_TCO与LCC测算核算模型.md",
    "07_附录与实务工具包/附录05_合同履约类_留样封样与质量索赔规范.md",
    "07_附录与实务工具包/附录06_工程采购类_工作范围SOW与现场签证审核.md",
    "07_附录与实务工具包/附录07_信息化采购类_选型打分表与SLA审查表.md",
    "07_附录与实务工具包/附录08_AI采购类_基准评测集与数据安全审查表.md",
    "07_附录与实务工具包/附录09_风险控制类_风险矩阵与应急保供预案.md",
    "07_附录与实务工具包/附录10_绩效评价类_平衡计分卡BSC指标字典.md"
]

doc = docx.Document()

# Set page margins (1 inch = 72pt)
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Cover Page
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(80)
title_p.paragraph_format.space_after = Pt(18)
run_title = title_p.add_run("化工行业采购管理")
run_title.font.name = "黑体"
run_title.font.size = Pt(32)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(26, 54, 93)
run_title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(40)
run_sub = sub_p.add_run("供应资源配置、业务实务与数智化跃迁\nProcurement Management in the Chemical Industry")
run_sub.font.name = "楷体"
run_sub.font.size = Pt(15)
run_sub.font.color.rgb = RGBColor(74, 85, 104)
run_sub._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")

meta_p = doc.add_paragraph()
meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta_p.paragraph_format.space_before = Pt(120)
meta_run = meta_p.add_run("专业专著 · 出版送审稿\n全六篇 · 十八章 · 十大实务工具包\n总字数：约 38 万字\n2026 年 10 月")
meta_run.font.name = "仿宋"
meta_run.font.size = Pt(12)
meta_run.font.color.rgb = RGBColor(113, 128, 150)
meta_run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")

doc.add_page_break()

def format_run(run, font_name="宋体", size_pt=10.5, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    xml_str = '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">' \
              '<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/>' \
              '<w:left w:val="none"/>' \
              '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/>' \
              '<w:right w:val="none"/>' \
              '</w:tcBorders>'
    tcBorders = parse_xml(xml_str)
    tcPr.append(tcBorders)

for rel_path in chapters_order:
    p = os.path.join(base_dir, rel_path.replace("/", os.sep))
    if not os.path.exists(p):
        continue
    with open(p, "r", encoding="utf-8") as fp:
        lines = fp.readlines()
    
    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []
    
    for line in lines:
        raw_line = line.rstrip("\r\n")
        
        # Code block
        if raw_line.strip().startswith("```"):
            if in_code_block:
                in_code_block = False
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.2)
                p_code.paragraph_format.space_before = Pt(3)
                p_code.paragraph_format.space_after = Pt(4)
                code_text = "\n".join(code_lines)
                r = p_code.add_run(code_text)
                format_run(r, font_name="Consolas", size_pt=8.5, color_rgb=RGBColor(45, 55, 72))
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue
        
        if in_code_block:
            code_lines.append(raw_line)
            continue
            
        # Table detection
        if raw_line.strip().startswith("|") and raw_line.strip().endswith("|"):
            if not in_table:
                in_table = True
                table_lines = [raw_line]
            else:
                table_lines.append(raw_line)
            continue
        else:
            if in_table:
                # Flush table
                in_table = False
                parsed_rows = []
                for tline in table_lines:
                    cells = [c.strip() for c in tline.strip("|").split("|")]
                    # skip separator line
                    if cells and all(re.match(r"^:?-+:?$", c) for c in cells):
                        continue
                    parsed_rows.append(cells)
                if parsed_rows:
                    num_cols = max(len(r) for r in parsed_rows)
                    norm_rows = []
                    for r in parsed_rows:
                        if len(r) < num_cols:
                            r.extend([""] * (num_cols - len(r)))
                        norm_rows.append(r[:num_cols])
                    
                    table = doc.add_table(rows=len(norm_rows), cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for i, row in enumerate(norm_rows):
                        for j, val in enumerate(row):
                            cell = table.cell(i, j)
                            cell.text = val
                            p_cell = cell.paragraphs[0]
                            p_cell.paragraph_format.space_before = Pt(2)
                            p_cell.paragraph_format.space_after = Pt(2)
                            p_cell.paragraph_format.line_spacing = 1.15
                            if i == 0:
                                p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                for r in p_cell.runs:
                                    format_run(r, font_name="黑体", size_pt=9.0, bold=True, color_rgb=RGBColor(26, 54, 93))
                                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EDF2F7"/>')
                                cell._tc.get_or_add_tcPr().append(shd)
                            else:
                                for r in p_cell.runs:
                                    format_run(r, font_name="宋体", size_pt=8.5, color_rgb=RGBColor(45, 55, 72))
                            set_cell_border(cell)
                    doc.add_paragraph()
                table_lines = []
        
        # Headers & Body
        s_line = raw_line.strip()
        if not s_line:
            continue
        if s_line == "---":
            continue
            
        if s_line.startswith("# "):
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(18)
            p_h.paragraph_format.space_after = Pt(8)
            r = p_h.add_run(s_line[2:].strip())
            format_run(r, font_name="黑体", size_pt=18, bold=True, color_rgb=RGBColor(26, 54, 93))
        elif s_line.startswith("## "):
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(13)
            p_h.paragraph_format.space_after = Pt(5)
            r = p_h.add_run(s_line[3:].strip())
            format_run(r, font_name="黑体", size_pt=14, bold=True, color_rgb=RGBColor(43, 108, 176))
        elif s_line.startswith("### "):
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(9)
            p_h.paragraph_format.space_after = Pt(4)
            r = p_h.add_run(s_line[4:].strip())
            format_run(r, font_name="黑体", size_pt=12, bold=True, color_rgb=RGBColor(45, 55, 72))
        elif s_line.startswith("#### "):
            p_h = doc.add_paragraph()
            p_h.paragraph_format.space_before = Pt(6)
            p_h.paragraph_format.space_after = Pt(3)
            r = p_h.add_run(s_line[5:].strip())
            format_run(r, font_name="黑体", size_pt=11, bold=True, color_rgb=RGBColor(74, 85, 104))
        elif s_line.startswith("> "):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Inches(0.25)
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(4)
            r = p_q.add_run(s_line[2:].strip())
            format_run(r, font_name="楷体", size_pt=10.5, italic=True, color_rgb=RGBColor(74, 85, 104))
        elif s_line.startswith("- ") or s_line.startswith("* "):
            p_l = doc.add_paragraph(style="List Bullet")
            p_l.paragraph_format.space_before = Pt(2)
            p_l.paragraph_format.space_after = Pt(2)
            p_l.paragraph_format.line_spacing = 1.25
            r = p_l.add_run(s_line[2:].strip())
            format_run(r, font_name="宋体", size_pt=10.5, color_rgb=RGBColor(45, 55, 72))
        elif re.match(r"^\d+\.\s", s_line):
            idx = s_line.find(".")
            p_l = doc.add_paragraph(style="List Number")
            p_l.paragraph_format.space_before = Pt(2)
            p_l.paragraph_format.space_after = Pt(2)
            p_l.paragraph_format.line_spacing = 1.25
            r = p_l.add_run(s_line[idx+1:].strip())
            format_run(r, font_name="宋体", size_pt=10.5, color_rgb=RGBColor(45, 55, 72))
        else:
            p_b = doc.add_paragraph()
            p_b.paragraph_format.first_line_indent = Inches(0.28)
            p_b.paragraph_format.space_before = Pt(2)
            p_b.paragraph_format.space_after = Pt(2)
            p_b.paragraph_format.line_spacing = 1.25
            r = p_b.add_run(s_line)
            format_run(r, font_name="宋体", size_pt=10.5, color_rgb=RGBColor(45, 55, 72))
    
    doc.add_page_break()

doc.save(output_docx)
print("Successfully generated Word document!")
print("Path:", output_docx)
print("Size:", os.path.getsize(output_docx), "bytes")
